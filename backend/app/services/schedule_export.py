"""Project schedule → Excel: activity table on the left, Gantt bars as filled cells on the right."""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.services.schedule_service import get_schedule

BRAND = "0066CC"
GROUP = "1A1A2E"
GROUP_ROW_BG = "F3F5F9"
HEADER_BG = "F1F3F6"
GRID = "E2E5EA"

# Daily columns stay readable up to ~4 months; beyond that switch to weeks.
DAILY_LIMIT_DAYS = 120

TABLE_COLS = [
    ("WBS", 8),
    ("Type", 9),
    ("Floor", 8),
    ("Name", 40),
    ("Days", 7),
    ("Start", 12),
    ("Finish", 12),
    ("Predecessor", 26),
    ("Lag", 6),
    ("Qty", 9),
    ("Unit", 7),
]


def _tree(activities: list[dict]) -> list[tuple[dict, int, str, bool]]:
    """Depth-first rows: (activity, depth, wbs, has_dates)."""
    ids = {a["id"] for a in activities}
    children: dict[int | None, list[dict]] = defaultdict(list)
    for a in activities:
        children[a["parent_id"] if a["parent_id"] in ids else None].append(a)

    dated_memo: dict[int, bool] = {}

    def dated(a: dict, guard: int = 0) -> bool:
        if a["row_type"] != "group":
            return True
        if guard > len(activities):
            return False
        if a["id"] not in dated_memo:
            dated_memo[a["id"]] = any(dated(k, guard + 1) for k in children.get(a["id"], []))
        return dated_memo[a["id"]]

    rows: list[tuple[dict, int, str, bool]] = []
    seen: set[int] = set()

    def walk(parent: int | None, depth: int, prefix: str) -> None:
        for n, a in enumerate(children.get(parent, []), start=1):
            if a["id"] in seen:
                continue
            seen.add(a["id"])
            wbs = f"{prefix}{n}"
            rows.append((a, depth, wbs, dated(a)))
            walk(a["id"], depth + 1, f"{wbs}.")

    walk(None, 0, "")
    return rows


def _d(iso: str | None) -> date | None:
    return date.fromisoformat(iso) if iso else None


def export_schedule(db: Session, project_id: int) -> tuple[bytes, str]:
    data = get_schedule(db, project_id)
    project = data["project"]
    activities = data["activities"]
    rows = _tree(activities)
    # COD is internal-only; users identify activities by name (+ floor).
    label_by_id = {
        a["id"]: f'{a["name"]} · {a["floor_name"]}' if a.get("floor_name") else a["name"]
        for a in activities
    }

    dated_rows = [r for r in rows if r[3]]
    if dated_rows:
        min_d = min(_d(r[0]["start_date"]) for r in dated_rows)
        max_d = max(_d(r[0]["finish_date"]) for r in dated_rows)
    else:
        min_d = max_d = date.today()
    span = (max_d - min_d).days + 1
    daily = span <= DAILY_LIMIT_DAYS
    if daily:
        buckets = [(min_d + timedelta(days=i), min_d + timedelta(days=i)) for i in range(span)]
    else:
        first = min_d - timedelta(days=min_d.weekday())  # Monday
        buckets = []
        cur = first
        while cur <= max_d:
            buckets.append((cur, cur + timedelta(days=6)))
            cur += timedelta(days=7)

    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"

    thin = Side(style="thin", color=GRID)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, color="555555", size=9)
    header_fill = PatternFill("solid", fgColor=HEADER_BG)
    bar_fill = PatternFill("solid", fgColor=BRAND)
    group_bar_fill = PatternFill("solid", fgColor=GROUP)
    group_row_fill = PatternFill("solid", fgColor=GROUP_ROW_BG)

    # Title
    ws["A1"] = f"{project['name']} — Project Schedule"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        f"Generated {datetime.now():%d/%m/%Y %H:%M} · "
        f"{sum(1 for r in rows if r[0]['row_type'] != 'group')} activities · "
        f"{min_d:%d/%m/%Y} – {max_d:%d/%m/%Y} · "
        f"timeline in {'days' if daily else 'weeks'}"
    )
    ws["A2"].font = Font(color="777777", size=9)

    month_row, head_row, first_data = 4, 5, 6
    first_tl_col = len(TABLE_COLS) + 1

    for c, (label, width) in enumerate(TABLE_COLS, start=1):
        cell = ws.cell(row=head_row, column=c, value=label)
        cell.font, cell.fill, cell.border = header_font, header_fill, border
        ws.column_dimensions[get_column_letter(c)].width = width
        ws.merge_cells(start_row=month_row, start_column=c, end_row=head_row, end_column=c)
        ws.cell(row=month_row, column=c).fill = header_fill
        cell.alignment = Alignment(vertical="center")
        ws.cell(row=month_row, column=c).value = label
        ws.cell(row=month_row, column=c).font = header_font
        ws.cell(row=month_row, column=c).alignment = Alignment(vertical="center")

    # Timeline headers: month band on top, day number / week start below
    month_start_col = first_tl_col
    for i, (b_start, _) in enumerate(buckets):
        col = first_tl_col + i
        ws.column_dimensions[get_column_letter(col)].width = 3.2 if daily else 4.5
        cell = ws.cell(row=head_row, column=col, value=b_start.day)
        cell.font = Font(size=8, color="777777")
        cell.fill, cell.border = header_fill, border
        cell.alignment = Alignment(horizontal="center")
        is_last = i == len(buckets) - 1
        next_month = None if is_last else buckets[i + 1][0].month
        if is_last or next_month != b_start.month:
            ws.merge_cells(start_row=month_row, start_column=month_start_col, end_row=month_row, end_column=col)
            m = ws.cell(row=month_row, column=month_start_col, value=f"{buckets[month_start_col - first_tl_col][0]:%b %Y}")
            m.font, m.fill = header_font, header_fill
            m.alignment = Alignment(horizontal="left")
            month_start_col = col + 1

    for r_i, (a, depth, wbs, has_dates) in enumerate(rows):
        r = first_data + r_i
        g = a["row_type"] == "group"
        start, finish = _d(a["start_date"]), _d(a["finish_date"])
        values = [
            wbs,
            "Group" if g else "Activity",
            a.get("floor_name") or "",
            a["name"],
            a["duration_days"] if has_dates else None,
            start if has_dates else None,
            finish if has_dates else None,
            label_by_id.get(a["predecessor_id"], "") if a["predecessor_id"] else "",
            a["lag_days"] if a["predecessor_id"] else None,
            a["quantity"],
            a["unit"] or "",
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            cell.font = Font(bold=g, size=10)
            if g:
                cell.fill = group_row_fill
            if isinstance(v, date):
                cell.number_format = "DD/MM/YYYY"
        name_col = next(i for i, (label, _) in enumerate(TABLE_COLS, start=1) if label == "Name")
        ws.cell(row=r, column=name_col).alignment = Alignment(indent=depth * 2)

        for i, (b_start, b_end) in enumerate(buckets):
            cell = ws.cell(row=r, column=first_tl_col + i)
            cell.border = border
            if has_dates and start <= b_end and finish >= b_start:
                cell.fill = group_bar_fill if g else bar_fill
            elif g:
                cell.fill = group_row_fill

    ws.freeze_panes = ws.cell(row=first_data, column=first_tl_col)
    ws.sheet_view.zoomScale = 90

    buf = BytesIO()
    wb.save(buf)
    safe = "".join(ch if ch.isalnum() else "_" for ch in project["name"]).strip("_")
    return buf.getvalue(), f"schedule_{safe or project_id}.xlsx"
